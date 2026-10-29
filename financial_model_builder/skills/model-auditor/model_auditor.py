#!/usr/bin/env python3
# SPDX-License-Identifier: MIT
# Copyright (c) 2026 Manigent Contributors
"""Audit an integrated financial model: 3-statement ties plus a workbook formula scan.

    python model_auditor.py ties  --csv model.csv [--map "bs_cash=Cash & equivalents"] [--tol 0.5]
    python model_auditor.py ties  --xlsx model.xlsx [--sheet IS --sheet BS --sheet CF] [--map ...]
    python model_auditor.py scan  --xlsx model.xlsx [--sheet DCF ...] [--allow-literals 0,1,12,100]
    python model_auditor.py all   --xlsx model.xlsx [...]          # ties + scan
    add --format table for a markdown findings table (default: JSON)

ties: reads statement rows (label in the first text column, one numeric column per
period, periods left to right) and checks every period:
    balance          total assets = total liabilities + equity (or = total L&E)
    cash_tie         CF ending cash = BS cash
    cf_sum           CF beginning cash + CFO + CFI + CFF = CF ending cash
    cash_continuity  this period's beginning cash = last period's ending cash
    re_roll          prior RE + net income - dividends = RE
    ppe_roll         prior PP&E + capex - D&A - disposals = PP&E
    debt_roll        prior debt + issuance - repayment = debt
    da_match         D&A on the IS = D&A add-back on the CF
    capex_sign       capex on the CF is negative
Rows are found by label aliases (see ALIASES); override with --map key=Label. When a
label appears twice (D&A on IS and CF, capex), the first is the IS/BS row and the
last is the CF row. Signs of dividends, capex, D&A and repayments are read as
magnitudes in roll-forwards.

scan (xlsx only, needs openpyxl): error values, formulas never calculated (openpyxl
does not calculate: recalc in Excel / LibreOffice first), numeric literals inside
formulas, typed numbers in rows that are otherwise formulas (pasted-over values),
formulas that break their row's relative pattern, external links, hidden sheets
and rows, circular references and whether iterative calculation is on.

CSV mode is stdlib only. Exit code 1 when any critical finding exists.
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import sys
from collections import Counter
from pathlib import Path

ERROR_VALUES = ("#REF!", "#VALUE!", "#DIV/0!", "#NAME?", "#N/A", "#NUM!", "#NULL!", "#SPILL!", "#CALC!")

# canonical key -> label aliases (compared after normalise()).
ALIASES: dict[str, tuple[str, ...]] = {
    "total_assets": ("total assets",),
    "total_liabilities": ("total liabilities",),
    "total_equity": ("total equity", "total stockholders equity", "total shareholders equity",
                     "total stockholders deficit", "total shareholders deficit", "total equity deficit"),
    "total_liabilities_and_equity": ("total liabilities and equity", "total liabilities and stockholders equity",
                                     "total liabilities and shareholders equity", "total liabilities equity"),
    "bs_cash": ("cash", "cash and cash equivalents", "cash and equivalents", "cash equivalents"),
    "cf_beginning_cash": ("beginning cash", "cash at beginning of period", "cash beginning of period",
                          "beginning cash balance", "cash at beginning of year", "opening cash"),
    "cf_ending_cash": ("ending cash", "cash at end of period", "cash end of period", "ending cash balance",
                       "cash at end of year", "closing cash"),
    "cfo": ("cfo", "cash from operations", "cash flow from operations", "cash from operating activities",
            "net cash from operating activities", "net cash provided by operating activities",
            "net cash provided by used in operating activities"),
    "cfi": ("cfi", "cash from investing", "cash flow from investing", "cash from investing activities",
            "net cash from investing activities", "net cash used in investing activities",
            "net cash provided by used in investing activities"),
    "cff": ("cff", "cash from financing", "cash flow from financing", "cash from financing activities",
            "net cash from financing activities", "net cash provided by used in financing activities"),
    "net_income": ("net income", "net income loss", "net earnings"),
    "retained_earnings": ("retained earnings", "retained earnings accumulated deficit", "accumulated deficit"),
    "dividends": ("dividends", "dividends paid", "dividends declared", "common dividends"),
    "d_and_a": ("d&a", "depreciation and amortization", "depreciation & amortization", "depreciation"),
    "ppe": ("pp&e, net", "net pp&e", "pp&e", "property, plant and equipment, net", "property and equipment, net"),
    "capex": ("capex", "capital expenditures", "purchases of pp&e", "purchase of pp&e",
              "purchases of property and equipment", "payments for property plant and equipment"),
    "disposals": ("disposals", "asset disposals", "pp&e disposals"),
    "debt": ("total debt", "debt"),
    "debt_issuance": ("debt issuance", "issuance of debt", "proceeds from debt", "proceeds from borrowings"),
    "debt_repayment": ("debt repayment", "repayment of debt", "repayments of debt", "mandatory amortization"),
}
TWICE = ("d_and_a", "capex")  # first occurrence = IS / BS row, last = CF row


def normalise(label: str) -> str:
    s = str(label).lower().replace("&", " and ").replace("'", "")
    s = re.sub(r"\(.*?\)", " ", s)
    s = re.sub(r"[^a-z0-9]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


_ALIAS_INDEX = {normalise(a): k for k, al in ALIASES.items() for a in al}


def to_number(v) -> float | None:
    if v is None:
        return None
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", "").replace("$", "")
    if not s or s in ("-", "—"):
        return None
    neg = s.startswith("(") and s.endswith(")")
    s = s.strip("()")
    try:
        x = float(s)
    except ValueError:
        return None
    return -x if neg else x


# ---------------------------------------------------------------- table extraction
def rows_from_csv(path: str) -> list[list]:
    with open(path, newline="", encoding="utf-8-sig") as fh:
        return [row for row in csv.reader(fh)]


def rows_from_xlsx(path: str, sheets: list[str] | None) -> tuple[list[list], list[str]]:
    """Cached values of the named sheets (or all), stacked in order, as plain rows.
    Also returns the sheet name of each row."""
    wb = _openpyxl().load_workbook(path, data_only=True)
    names = sheets or wb.sheetnames
    rows, origin = [], []
    for n in names:
        if n not in wb.sheetnames:
            raise SystemExit(f"sheet {n!r} not in {wb.sheetnames}")
        for r in wb[n].iter_rows(values_only=True):
            rows.append(list(r))
            origin.append(n)
    return rows, origin


def extract(rows: list[list], overrides: dict[str, str] | None = None) -> tuple[list[str], dict[str, list[float | None]]]:
    """(period labels, {canonical key: values}) from a statement grid.

    The label column is the first column holding text in most non-empty rows; period
    columns are those to its right holding numbers in at least two rows."""
    overrides = {k: normalise(v) for k, v in (overrides or {}).items()}
    width = max((len(r) for r in rows), default=0)
    text_counts = [sum(1 for r in rows if c < len(r) and isinstance(r[c], str) and to_number(r[c]) is None
                       and r[c].strip()) for c in range(width)]
    label_col = max(range(width), key=lambda c: (text_counts[c], -c)) if width else 0
    num_cols = [c for c in range(label_col + 1, width)
                if sum(1 for r in rows if c < len(r) and to_number(r[c]) is not None) >= 2]
    header = next((r for r in rows if label_col < len(r) and sum(
        1 for c in num_cols if c < len(r) and r[c] not in (None, "")) >= max(1, len(num_cols) // 2)
        and all(to_number(r[c]) is None or _looks_like_year(r[c]) for c in num_cols if c < len(r)
                and r[c] not in (None, ""))), None)
    periods = [str(header[c]) if header and c < len(header) and header[c] not in (None, "") else f"col{c + 1}"
               for c in num_cols]
    found: dict[str, list[list[float | None]]] = {}
    for r in rows:
        if label_col >= len(r) or not isinstance(r[label_col], str):
            continue
        lab = normalise(r[label_col])
        key = next((k for k, v in overrides.items() if v == lab), None) or _ALIAS_INDEX.get(lab)
        if not key:
            continue
        vals = [to_number(r[c]) if c < len(r) else None for c in num_cols]
        if r is header:
            continue
        found.setdefault(key, []).append(vals)
    table: dict[str, list[float | None]] = {}
    for key, occ in found.items():
        if key in TWICE and len(occ) >= 2:
            table[("is_" if key == "d_and_a" else "bs_") + key] = occ[0]
            table["cf_" + key] = occ[-1]
        elif key in TWICE:
            table["cf_" + key if key == "capex" else "is_" + key] = occ[0]
        else:
            table[key] = occ[0]
    return periods, table


def _looks_like_year(v) -> bool:
    x = to_number(v)
    return x is not None and float(x).is_integer() and 1900 <= x <= 2200


# ---------------------------------------------------------------- tie checks
def _f(sev, check, period, detail, diff=None, sheet="", cell=""):
    return {"severity": sev, "check": check, "sheet": sheet, "cell": cell, "period": period,
            "detail": detail, "difference": None if diff is None else round(diff, 6)}


def tie_checks(periods: list[str], t: dict[str, list[float | None]], tol: float = 0.5) -> list[dict]:
    out: list[dict] = []
    n = len(periods)

    def v(key, i):
        arr = t.get(key)
        return None if arr is None or i < 0 or i >= len(arr) else arr[i]

    def off(x):
        return abs(x) > tol

    ran = Counter()
    for i, p in enumerate(periods):
        a, l, e, tle = v("total_assets", i), v("total_liabilities", i), v("total_equity", i), v("total_liabilities_and_equity", i)
        if a is not None and (tle is not None or (l is not None and e is not None)):
            ran["balance"] += 1
            rhs = tle if tle is not None else l + e
            if off(a - rhs):
                out.append(_f("critical", "balance", p, "total assets != liabilities + equity; fix this first, "
                              "everything downstream is suspect", a - rhs))
        ce, bc = v("cf_ending_cash", i), v("bs_cash", i)
        if ce is not None and bc is not None:
            ran["cash_tie"] += 1
            if off(ce - bc):
                out.append(_f("critical", "cash_tie", p, "CF ending cash != BS cash", ce - bc))
        cb, o, inv, fin = v("cf_beginning_cash", i), v("cfo", i), v("cfi", i), v("cff", i)
        if None not in (cb, o, inv, fin, ce):
            ran["cf_sum"] += 1
            d = cb + o + inv + fin - ce
            if off(d):
                out.append(_f("critical", "cf_sum", p, "beginning cash + CFO + CFI + CFF != ending cash", d))
        if i > 0 and cb is not None and v("cf_ending_cash", i - 1) is not None:
            ran["cash_continuity"] += 1
            d = cb - v("cf_ending_cash", i - 1)
            if off(d):
                out.append(_f("critical", "cash_continuity", p, "beginning cash != prior period ending cash", d))
        re_p, re_c, ni = v("retained_earnings", i - 1), v("retained_earnings", i), v("net_income", i)
        if i > 0 and None not in (re_p, re_c, ni):
            ran["re_roll"] += 1
            d = re_p + ni - abs(v("dividends", i) or 0.0) - re_c
            if off(d):
                out.append(_f("critical", "re_roll", p, "prior RE + net income - dividends != RE "
                              "(other equity items such as buybacks retired against RE explain a gap)", d))
        ppe_p, ppe_c = v("ppe", i - 1), v("ppe", i)
        capex = v("cf_capex", i) if v("cf_capex", i) is not None else v("bs_capex", i)
        da = v("cf_d_and_a", i) if v("cf_d_and_a", i) is not None else v("is_d_and_a", i)
        if i > 0 and None not in (ppe_p, ppe_c, capex, da):
            ran["ppe_roll"] += 1
            d = ppe_p + abs(capex) - abs(da) - abs(v("disposals", i) or 0.0) - ppe_c
            if off(d):
                out.append(_f("warning", "ppe_roll", p, "prior PP&E + capex - D&A - disposals != PP&E "
                              "(D&A that includes intangible amortization breaks this; split it)", d))
        dp, dc = v("debt", i - 1), v("debt", i)
        if i > 0 and None not in (dp, dc) and (v("debt_issuance", i) is not None or v("debt_repayment", i) is not None):
            ran["debt_roll"] += 1
            d = dp + abs(v("debt_issuance", i) or 0.0) - abs(v("debt_repayment", i) or 0.0) - dc
            if off(d):
                out.append(_f("critical", "debt_roll", p, "prior debt + issuance - repayment != debt", d))
        isda, cfda = v("is_d_and_a", i), v("cf_d_and_a", i)
        if isda is not None and cfda is not None:
            ran["da_match"] += 1
            d = abs(isda) - abs(cfda)
            if off(d):
                out.append(_f("warning", "da_match", p, "D&A on the IS != D&A add-back on the CF", d))
        cfc = v("cf_capex", i)
        if cfc is not None:
            ran["capex_sign"] += 1
            if cfc > tol:
                out.append(_f("warning", "capex_sign", p, "capex on the CF is positive; cash outflows should be negative"))
    missing = [c for c in ("balance", "cash_tie", "cf_sum") if not ran[c]]
    if missing:
        out.append(_f("info", "coverage", "", f"not run (rows not found): {', '.join(missing)}; "
                      "pass --map key=Label for the rows the aliases missed"))
    out.append(_f("info", "coverage", "", "ran: " + ", ".join(f"{k} x{c}" for k, c in sorted(ran.items()))))
    return out


# ---------------------------------------------------------------- formula scan
_REF_FULL = re.compile(
    r"(?<![A-Za-z0-9_.])((?:'[^']+'|[A-Za-z0-9_.]+)!)?(\$?)([A-Z]{1,3})(\$?)([0-9]{1,7})"
    r"(?::(\$?)([A-Z]{1,3})(\$?)([0-9]{1,7}))?(?![0-9(A-Za-z_])")
_STRING = re.compile(r'"[^"]*"')
_NUM = re.compile(r"(?<![A-Za-z0-9_.$:!])(\d+\.?\d*(?:[eE][+-]?\d+)?|\.\d+)(?![A-Za-z0-9_(:!])")
_EXTERNAL = re.compile(r"\[[^\]]+\]|\.xls[xmb]?['!]", re.I)


def col_to_num(col: str) -> int:
    n = 0
    for ch in col:
        n = n * 26 + ord(ch) - 64
    return n


def num_to_col(n: int) -> str:
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def relative_pattern(formula: str, row: int, col: int) -> str:
    """Formula with each A1 reference rewritten relative to (row, col), R1C1-style."""
    def rep(m):
        sheet, cabs, c, rabs, r = m.group(1) or "", m.group(2), m.group(3), m.group(4), m.group(5)
        part = (f"C{col_to_num(c)}" if cabs else f"C[{col_to_num(c) - col}]") + \
               (f"R{r}" if rabs else f"R[{int(r) - row}]")
        if m.group(7):
            c2, r2 = m.group(7), m.group(9)
            part += ":" + ((f"C{col_to_num(c2)}" if m.group(6) else f"C[{col_to_num(c2) - col}]") +
                           (f"R{r2}" if m.group(8) else f"R[{int(r2) - row}]"))
        return sheet + part
    return _REF_FULL.sub(rep, _STRING.sub('""', formula))


def literals_in(formula: str) -> list[str]:
    body = _STRING.sub('""', formula.lstrip("="))
    body = _REF_FULL.sub("REF", body)
    return _NUM.findall(body)


def refs_in(formula: str, sheet: str, max_range: int = 5000) -> list[tuple[str, str]]:
    """(sheet, A1) cells a formula reads; ranges expanded up to max_range cells."""
    out = []
    for m in _REF_FULL.finditer(_STRING.sub('""', formula)):
        sh = (m.group(1) or "").rstrip("!").strip("'") or sheet
        c1, r1 = col_to_num(m.group(3)), int(m.group(5))
        if m.group(7):
            c2, r2 = col_to_num(m.group(7)), int(m.group(9))
            if (abs(c2 - c1) + 1) * (abs(r2 - r1) + 1) > max_range:
                continue
            for c in range(min(c1, c2), max(c1, c2) + 1):
                for r in range(min(r1, r2), max(r1, r2) + 1):
                    out.append((sh, f"{num_to_col(c)}{r}"))
        else:
            out.append((sh, f"{num_to_col(c1)}{r1}"))
    return out


def find_cycles(graph: dict[tuple, list[tuple]], limit: int = 20) -> list[list[tuple]]:
    """Strongly connected components with more than one node, or a self-loop
    (iterative Tarjan, so large models do not hit the recursion limit)."""
    index: dict = {}
    low: dict = {}
    on: set = set()
    stack: list = []
    out: list[list[tuple]] = []
    counter = 0
    for root in graph:
        if root in index:
            continue
        work = [(root, 0)]
        while work:
            v, i = work.pop()
            if i == 0:
                index[v] = low[v] = counter
                counter += 1
                stack.append(v)
                on.add(v)
            succ = [w for w in graph.get(v, ()) if w in graph]
            recurse = False
            for j in range(i, len(succ)):
                w = succ[j]
                if w not in index:
                    work.append((v, j + 1))
                    work.append((w, 0))
                    recurse = True
                    break
                if w in on:
                    low[v] = min(low[v], index[w])
            if recurse:
                continue
            if low[v] == index[v]:
                comp = []
                while True:
                    w = stack.pop()
                    on.discard(w)
                    comp.append(w)
                    if w == v:
                        break
                if len(comp) > 1 or v in graph.get(v, ()):
                    out.append(sorted(comp))
                    if len(out) >= limit:
                        return out
            if work:
                parent = work[-1][0]
                low[parent] = min(low[parent], low[v])
    return out


def scan_xlsx(path: str, sheets: list[str] | None = None, allow: set[float] | None = None) -> list[dict]:
    opx = _openpyxl()
    allow = allow if allow is not None else {0.0, 1.0}
    wbf = opx.load_workbook(path, data_only=False)
    wbv = opx.load_workbook(path, data_only=True)
    out: list[dict] = []
    names = sheets or wbf.sheetnames
    graph: dict[tuple, list[tuple]] = {}
    uncalculated = 0
    calc = getattr(wbf, "calculation", None)
    iterate = bool(getattr(calc, "iterate", False))
    for ws in wbf.worksheets:
        if ws.sheet_state != "visible":
            out.append(_f("warning", "hidden_sheet", "", f"sheet is {ws.sheet_state}; it may hold overrides", sheet=ws.title))
    for name in names:
        ws, wv = wbf[name], wbv[name]
        hidden_rows = [r for r, d in ws.row_dimensions.items() if d.hidden]
        if hidden_rows:
            out.append(_f("info", "hidden_rows", "", f"hidden rows {hidden_rows[:20]}", sheet=name))
        for row in ws.iter_rows():
            formulas, constants, patterns = [], [], []
            for cell in row:
                val = cell.value
                cached = wv[cell.coordinate].value
                if isinstance(cached, str) and cached.strip() in ERROR_VALUES:
                    out.append(_f("critical", "error_value", "", f"{cached.strip()}", sheet=name, cell=cell.coordinate))
                if isinstance(val, str) and val.startswith("="):
                    formulas.append(cell)
                    if cached is None:
                        uncalculated += 1
                    if _EXTERNAL.search(val):
                        out.append(_f("warning", "external_link", "", val, sheet=name, cell=cell.coordinate))
                    lits = [x for x in literals_in(val) if float(x) not in allow]
                    if lits:
                        out.append(_f("warning", "hardcode_in_formula", "", f"{val}  (literals: {', '.join(lits)})",
                                      sheet=name, cell=cell.coordinate))
                    patterns.append((cell, relative_pattern(val, cell.row, cell.column)))
                    graph[(name, cell.coordinate)] = refs_in(val, name)
                elif isinstance(val, (int, float)) and not isinstance(val, bool):
                    constants.append(cell)
            if len(formulas) >= 3 and constants:
                first_formula_col = min(c.column for c in formulas)
                for c in constants:
                    if c.column > first_formula_col:
                        out.append(_f("warning", "pasted_value", "", f"typed number {c.value} in a row of formulas",
                                      sheet=name, cell=c.coordinate))
            if len(patterns) >= 3:
                common, count = Counter(p for _, p in patterns).most_common(1)[0]
                if count >= 2:
                    for cell, p in patterns:
                        if p != common:
                            out.append(_f("warning", "inconsistent_formula", "",
                                          f"{cell.value} breaks the row pattern "
                                          f"(first projection column often links differently; check it)",
                                          sheet=name, cell=cell.coordinate))
    if uncalculated:
        out.append(_f("warning", "not_calculated", "", f"{uncalculated} formula cells have no cached value: the file "
                      "was written by a library and never recalculated. Recalculate (LibreOffice headless or Excel) "
                      "before trusting ties or error checks"))
    cycles = find_cycles(graph)
    for comp in cycles:
        cells = ", ".join(f"{s}!{c}" for s, c in comp[:8]) + (" ..." if len(comp) > 8 else "")
        sev = "info" if iterate else "critical"
        out.append(_f(sev, "circular_reference", "", f"loop of {len(comp)} cells: {cells}. Iterative calculation is "
                      f"{'ON' if iterate else 'OFF'}; document it on the cover and keep a switch that breaks the loop"))
    out.append(_f("info", "iterative_calc", "", f"iterate={iterate} count={getattr(calc, 'iterateCount', None)} "
                  f"delta={getattr(calc, 'iterateDelta', None)}"))
    return out


def _openpyxl():
    try:
        import openpyxl  # noqa: PLC0415
    except ImportError:
        raise SystemExit("openpyxl is needed for .xlsx (pip install openpyxl); CSV mode needs nothing")
    return openpyxl


# ---------------------------------------------------------------- CLI
def summarise(findings: list[dict]) -> dict:
    c = Counter(f["severity"] for f in findings)
    overall = "Major issues" if c["critical"] else ("Minor issues" if c["warning"] else "Clean")
    return {"overall": overall, "critical": c["critical"], "warning": c["warning"], "info": c["info"]}


def as_table(findings: list[dict]) -> str:
    s = summarise(findings)
    lines = [f"Overall: {s['overall']} - {s['critical']} critical, {s['warning']} warnings, {s['info']} info", "",
             "| # | Severity | Check | Sheet | Cell | Period | Detail | Difference |", "|---|---|---|---|---|---|---|---|"]
    order = {"critical": 0, "warning": 1, "info": 2}
    for i, f in enumerate(sorted(findings, key=lambda f: order[f["severity"]]), 1):
        detail = str(f["detail"]).replace("|", "\\|")
        lines.append(f"| {i} | {f['severity']} | {f['check']} | {f['sheet']} | {f['cell']} | {f['period']} | "
                     f"{detail} | {'' if f['difference'] is None else f['difference']} |")
    return "\n".join(lines)


def parse_map(items: list[str] | None) -> dict[str, str]:
    out = {}
    for it in items or []:
        for part in it.split(","):
            if "=" in part:
                k, v = part.split("=", 1)
                out[k.strip()] = v.strip()
    return out


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("mode", choices=("ties", "scan", "all"))
    p.add_argument("--csv")
    p.add_argument("--xlsx")
    p.add_argument("--sheet", action="append")
    p.add_argument("--map", action="append", help="key=Row label (repeatable or comma-separated)")
    p.add_argument("--tol", type=float, default=0.5, help="absolute tolerance (model units)")
    p.add_argument("--allow-literals", default="0,1", help="numbers allowed inside formulas")
    p.add_argument("--format", choices=("json", "table"), default="json")
    a = p.parse_args(argv)
    if not (a.csv or a.xlsx):
        p.error("--csv or --xlsx is required")
    if a.mode in ("scan", "all") and not a.xlsx:
        p.error("scan needs --xlsx")
    findings: list[dict] = []
    if a.mode in ("ties", "all"):
        rows = rows_from_csv(a.csv) if a.csv else rows_from_xlsx(a.xlsx, a.sheet)[0]
        periods, table = extract(rows, parse_map(a.map))
        findings += tie_checks(periods, table, a.tol)
    if a.mode in ("scan", "all"):
        allow = {float(x) for x in a.allow_literals.split(",") if x.strip()}
        findings += scan_xlsx(a.xlsx, a.sheet, allow)
    if a.format == "table":
        print(as_table(findings))
    else:
        print(json.dumps({"summary": summarise(findings), "findings": findings}, indent=1))
    return 1 if summarise(findings)["critical"] else 0


if __name__ == "__main__":
    sys.exit(main())
