# SPDX-License-Identifier: MIT
# Copyright (c) 2026 Manigent Contributors
"""Unit tests for model_auditor.py on synthetic statements and workbooks.

Run: python3 -m pytest test_model_auditor.py   (xlsx tests skip without openpyxl)
"""
import csv
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import model_auditor as ma  # noqa: E402

try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None

# A tiny synthetic company that ties in every period.
CLEAN = [
    ["Synthetic Co", "", "", ""],
    ["($ thousands)", "2024A", "2025E", "2026E"],
    ["Revenue", 100, 110, 121],
    ["Depreciation & amortization", 10, 11, 12],
    ["Net income", 20, 22, 24],
    ["Cash and cash equivalents", 50, 66, 84],
    ["PP&E, net", 200, 204, 210],
    ["Total assets", 250, 270, 294],
    ["Total liabilities", 150, 150, 150],
    ["Retained earnings", 60, 80, 102],
    ["Total equity", 100, 120, 144],
    ["Total liabilities & equity", 250, 270, 294],
    ["Beginning cash", 40, 50, 66],
    ["Depreciation & amortization", 10, 11, 12],
    ["Net cash from operating activities", 30, 33, 36],
    ["Capital expenditures", -15, -15, -18],
    ["Net cash from investing activities", -15, -15, -18],
    ["Dividends paid", -5, -2, -2],
    ["Net cash from financing activities", -5, -2, 0],
    ["Ending cash", 50, 66, 84],
]


def write_csv(rows, d):
    p = Path(d) / "m.csv"
    with open(p, "w", newline="") as fh:
        csv.writer(fh).writerows(rows)
    return str(p)


def checks(findings, sev=None):
    return [f["check"] for f in findings if f["severity"] != "info" and (sev is None or f["severity"] == sev)]


class Extract(unittest.TestCase):
    def test_labels_periods_and_twice_rows(self):
        periods, t = ma.extract(CLEAN)
        self.assertEqual(periods, ["2024A", "2025E", "2026E"])
        self.assertEqual(t["is_d_and_a"], [10, 11, 12])
        self.assertEqual(t["cf_d_and_a"], [10, 11, 12])
        self.assertEqual(t["cf_capex"], [-15, -15, -18])
        self.assertIn("total_liabilities_and_equity", t)

    def test_number_parsing(self):
        self.assertEqual(ma.to_number("(1,234.5)"), -1234.5)
        self.assertEqual(ma.to_number("$12"), 12)
        self.assertIsNone(ma.to_number("-"))

    def test_map_override(self):
        rows = [r[:] for r in CLEAN]
        rows[5][0] = "Liquidity bucket"
        _, t = ma.extract(rows, {"bs_cash": "Liquidity bucket"})
        self.assertEqual(t["bs_cash"], [50, 66, 84])


class Ties(unittest.TestCase):
    def run_ties(self, rows):
        periods, t = ma.extract(rows)
        return ma.tie_checks(periods, t, tol=0.5)

    def test_clean_model_has_no_findings(self):
        f = self.run_ties(CLEAN)
        self.assertEqual(checks(f), [], f)
        ran = next(x for x in f if x["check"] == "coverage" and x["detail"].startswith("ran"))
        for name in ("balance", "cash_tie", "cf_sum", "cash_continuity", "re_roll", "ppe_roll", "da_match", "capex_sign"):
            self.assertIn(name, ran["detail"])

    def test_unbalanced_period_is_critical(self):
        rows = [r[:] for r in CLEAN]
        rows[7][3] = 300  # total assets 2026E
        f = self.run_ties(rows)
        bal = [x for x in f if x["check"] == "balance"]
        self.assertEqual(len(bal), 1)
        self.assertEqual((bal[0]["period"], bal[0]["severity"], bal[0]["difference"]), ("2026E", "critical", 6))

    def test_cash_breaks(self):
        rows = [r[:] for r in CLEAN]
        rows[19][2] = 70   # ending cash 2025E no longer = BS cash, = beginning + flows, = next beginning
        f = self.run_ties(rows)
        self.assertIn("cash_tie", checks(f, "critical"))
        self.assertIn("cf_sum", checks(f, "critical"))
        self.assertIn("cash_continuity", checks(f, "critical"))

    def test_roll_forwards_and_signs(self):
        rows = [r[:] for r in CLEAN]
        rows[9][2] = 81      # RE 2025E off by one
        rows[6][3] = 215     # PP&E 2026E off
        rows[13][1] = 9      # CF D&A differs from IS D&A in 2024A
        rows[15][2] = 15     # positive capex on CF
        f = self.run_ties(rows)
        names = checks(f)
        for n in ("re_roll", "ppe_roll", "da_match", "capex_sign"):
            self.assertIn(n, names)

    def test_missing_rows_reported(self):
        f = self.run_ties([["Item", "2025"], ["Revenue", 1], ["Revenue", 2]])
        self.assertTrue(any(x["check"] == "coverage" and "not run" in x["detail"] for x in f))

    def test_cli_csv_exit_code(self):
        with tempfile.TemporaryDirectory() as d:
            self.assertEqual(ma.main(["ties", "--csv", write_csv(CLEAN, d)]), 0)
            bad = [r[:] for r in CLEAN]
            bad[7][1] = 251
            self.assertEqual(ma.main(["ties", "--csv", write_csv(bad, d), "--format", "table"]), 1)


class Formulas(unittest.TestCase):
    def test_relative_pattern_matches_across_a_row(self):
        a = ma.relative_pattern("=C5*(1+$B$2)", 5, 4)
        b = ma.relative_pattern("=D5*(1+$B$2)", 5, 5)
        self.assertEqual(a, b)
        self.assertNotEqual(a, ma.relative_pattern("=D5*(1+$B$3)", 5, 5))

    def test_literals(self):
        self.assertEqual(ma.literals_in("=C5*1.05"), ["1.05"])
        self.assertEqual(ma.literals_in('=SUM(C1:C9)+LOG10(A1)'), [])
        self.assertEqual(ma.literals_in('="FY 2025"&A1'), [])
        self.assertEqual(ma.literals_in("=Inputs!B3*Sheet2!C4"), [])

    def test_cycle_finder(self):
        g = {("S", "A1"): [("S", "B1")], ("S", "B1"): [("S", "C1")], ("S", "C1"): [("S", "A1")],
             ("S", "D1"): [("S", "A1")], ("S", "E1"): [("S", "E1")]}
        cycles = ma.find_cycles(g)
        self.assertEqual(sorted(len(c) for c in cycles), [1, 3])

    def test_cycle_finder_handles_long_chains(self):
        g = {("S", f"A{i}"): [("S", f"A{i + 1}")] for i in range(1, 20000)}
        g[("S", "A20000")] = [("S", "A1")]
        self.assertEqual(len(ma.find_cycles(g)[0]), 20000)


@unittest.skipIf(openpyxl is None, "openpyxl not installed")
class Scan(unittest.TestCase):
    def build(self, d, iterate=False):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Model"
        ws.append(["Line", "2024A", "2025E", "2026E", "2027E", "2028E"])
        ws.append(["Growth", None, 0.1])
        ws.append(["Revenue", 100, "=B3*(1+$C$2)", "=C3*(1+$C$2)", "=D3*1.05", "=E3*(1+$C$2)"])  # hard-code in E3
        ws.append(["Costs", 60, "=C3*0.6", "=D3*0.6", 70, "=F3*0.6"])                        # pasted value E4
        ws.append(["Link", None, "=[1]Other!A1", None, None, None])                          # external link
        ws.append(["Interest", None, "=C7*0.05", None, None, None])
        ws.append(["Debt", None, "=100-C6", None, None, None])                               # loop C6 <-> C7
        ws["B8"] = "#REF!"
        hidden = wb.create_sheet("Overrides")
        hidden.sheet_state = "hidden"
        wb.calculation.iterate = iterate
        p = Path(d) / "m.xlsx"
        wb.save(p)
        return str(p)

    def test_scan_finds_planted_issues(self):
        with tempfile.TemporaryDirectory() as d:
            f = ma.scan_xlsx(self.build(d), allow={0.0, 1.0})
        by = {}
        for x in f:
            by.setdefault(x["check"], []).append(x)
        self.assertIn("E3", {x["cell"] for x in by["hardcode_in_formula"]})
        self.assertEqual([x["cell"] for x in by["pasted_value"]], ["E4"])
        self.assertIn("E3", {x["cell"] for x in by["inconsistent_formula"]})
        self.assertEqual([x["cell"] for x in by["external_link"]], ["C5"])
        self.assertEqual(by["error_value"][0]["cell"], "B8")
        self.assertEqual(by["hidden_sheet"][0]["sheet"], "Overrides")
        self.assertEqual(by["circular_reference"][0]["severity"], "critical")
        self.assertIn("not_calculated", by)

    def test_circularity_with_iteration_on_is_info(self):
        with tempfile.TemporaryDirectory() as d:
            f = ma.scan_xlsx(self.build(d, iterate=True))
        circ = [x for x in f if x["check"] == "circular_reference"]
        self.assertEqual(circ[0]["severity"], "info")

    def test_ties_from_xlsx(self):
        with tempfile.TemporaryDirectory() as d:
            wb = openpyxl.Workbook()
            for r in CLEAN:
                wb.active.append(r)
            p = Path(d) / "s.xlsx"
            wb.save(p)
            self.assertEqual(ma.main(["ties", "--xlsx", str(p)]), 0)


if __name__ == "__main__":
    unittest.main()
