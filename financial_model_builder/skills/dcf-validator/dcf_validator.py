#!/usr/bin/env python3
# SPDX-License-Identifier: MIT
# Copyright (c) 2026 Manigent Contributors
"""DCF-specific sanity checks on a workbook (cached values) or on a JSON of inputs.

    python dcf_validator.py --xlsx model.xlsx [--sheet DCF --sheet WACC] [--riskfree 0.0xx]
                            [--tv-method gordon|exit] [--format table]
    python dcf_validator.py --json inputs.json [--format table]

Inputs are located by row label (first number to the right of the label, any sheet):
WACC / discount rate, terminal (perpetuity) growth, riskfree, terminal value,
PV of terminal value, enterprise value, exit multiple, a mid-year switch, and the
discount-period row. JSON keys: wacc, terminal_growth, riskfree, terminal_value,
pv_terminal_value, enterprise_value, exit_multiple, mid_year (bool), discount_periods
(list), tv_method ('gordon' | 'exit').

Checks: formula error values; uncalculated formulas; g >= WACC (critical); g above the
riskfree rate; WACC outside a plausible band; percentages stored as whole numbers;
PV(TV) share of EV; discount periods vs the mid-year switch; the exponent actually used
to discount TV (from TV, PV(TV) and WACC) vs the end of year N or N - 0.5, and an exit-
multiple TV discounted at a mid-year exponent. Reports which inputs it could not find.

xlsx needs openpyxl; JSON mode is stdlib only. Exit code 1 when a critical finding exists.
"""
from __future__ import annotations

import argparse
import json
import math
import re
import sys

ERROR_VALUES = ("#REF!", "#VALUE!", "#DIV/0!", "#NAME?", "#N/A", "#NUM!", "#NULL!", "#SPILL!", "#CALC!")
WACC_BAND = (0.04, 0.20)
TV_SHARE_FLAG = 0.75

# Order matters: the first pattern that matches a label wins.
LABELS: list[tuple[str, re.Pattern]] = [
    ("pv_terminal_value", re.compile(r"(pv|present value|discounted).*(terminal|\btv\b)|(terminal value|\btv\b).*\b(pv|present value|discounted)\b")),
    ("exit_multiple", re.compile(r"exit multiple|terminal multiple|exit (ev ?/ ?)?ebitda")),
    ("terminal_growth", re.compile(r"terminal.*growth|perpetuity growth|perpetual growth|long[- ]term growth|\btgr\b")),
    ("terminal_value", re.compile(r"^(terminal value|tv)\b")),
    ("wacc", re.compile(r"\bwacc\b|weighted average cost of capital|^discount rate")),
    ("riskfree", re.compile(r"risk[- ]?free|\brf\b|10[- ]?y(ea)?r?\.? treasury|treasury yield")),
    ("enterprise_value", re.compile(r"enterprise value|^ev$|^implied ev")),
    ("mid_year", re.compile(r"mid[- ]?year")),
    ("discount_periods", re.compile(r"discount(ing)? period|period for discount|discount year|years? discounted")),
]


def to_number(v):
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", "")
    pct = s.endswith("%")
    try:
        x = float(s.rstrip("%"))
    except ValueError:
        return None
    return x / 100.0 if pct else x


def to_flag(v):
    if isinstance(v, bool):
        return v
    if isinstance(v, (int, float)):
        return bool(v)
    s = str(v or "").strip().lower()
    if s in ("yes", "y", "true", "on", "mid", "mid-year", "1"):
        return True
    if s in ("no", "n", "false", "off", "end", "end-year", "end of year", "0"):
        return False
    return None


def _f(sev, check, detail, where=""):
    return {"severity": sev, "check": check, "where": where, "detail": detail}


def read_xlsx(path: str, sheets: list[str] | None) -> tuple[dict, list[dict]]:
    try:
        import openpyxl  # noqa: PLC0415
    except ImportError:
        raise SystemExit("openpyxl is needed for .xlsx (pip install openpyxl), or use --json")
    wbv = openpyxl.load_workbook(path, data_only=True)
    wbf = openpyxl.load_workbook(path, data_only=False)
    names = sheets or wbv.sheetnames
    found: dict = {}
    findings: list[dict] = []
    uncalculated = 0
    for n in names:
        if n not in wbv.sheetnames:
            raise SystemExit(f"sheet {n!r} not in {wbv.sheetnames}")
        wv, wf = wbv[n], wbf[n]
        for row in wv.iter_rows():
            for cell in row:
                v = cell.value
                fv = wf[cell.coordinate].value
                if isinstance(fv, str) and fv.startswith("=") and v is None:
                    uncalculated += 1
                if isinstance(v, str) and v.strip() in ERROR_VALUES:
                    findings.append(_f("critical", "error_value", v.strip(), f"{n}!{cell.coordinate}"))
                    continue
                if not isinstance(v, str) or to_number(v) is not None:
                    continue
                label = v.strip().lower()
                key = next((k for k, pat in LABELS if pat.search(label)), None)
                if key is None or (key in found and key != "discount_periods"):
                    continue
                right = [c.value for c in row if c.column > cell.column][:24]
                where = f"{n}!{cell.coordinate}"
                if key == "discount_periods":
                    nums = [to_number(x) for x in right if to_number(x) is not None]
                    if nums and key not in found:
                        found[key] = (nums, where)
                elif key == "mid_year":
                    flag = next((to_flag(x) for x in right[:6] if to_flag(x) is not None), None)
                    if flag is not None:
                        found[key] = (flag, where)
                else:
                    num = next((to_number(x) for x in right[:8] if to_number(x) is not None), None)
                    if num is not None:
                        found[key] = (num, where)
    if uncalculated:
        findings.append(_f("warning", "not_calculated", f"{uncalculated} formula cells have no cached value; "
                           "recalculate (Excel or LibreOffice headless) before validating"))
    return found, findings


def validate(found: dict, riskfree: float | None = None, tv_method: str | None = None) -> list[dict]:
    """found: {key: (value, where)}. Returns findings."""
    out: list[dict] = []

    def get(k):
        return found[k][0] if k in found else None

    def where(k):
        return found[k][1] if k in found else ""

    for k in ("wacc", "terminal_growth", "riskfree"):
        v = get(k)
        if v is not None and abs(v) >= 1:
            out.append(_f("warning", "percent_as_whole_number", f"{k} = {v}; read as {v / 100:.4%}. Store rates "
                          "as decimals so formulas do not silently scale by 100", where(k)))
            found[k] = (v / 100.0, where(k))
    wacc, g = get("wacc"), get("terminal_growth")
    rf = riskfree if riskfree is not None else get("riskfree")
    if wacc is not None and g is not None:
        if g >= wacc:
            out.append(_f("critical", "growth_vs_wacc", f"terminal growth {g:.4%} >= WACC {wacc:.4%}: Gordon TV is "
                          "undefined or negative", where("terminal_growth")))
        else:
            out.append(_f("info", "growth_vs_wacc", f"terminal growth {g:.4%} < WACC {wacc:.4%}"))
    if g is not None and rf is not None and g > rf:
        out.append(_f("warning", "growth_vs_riskfree", f"terminal growth {g:.4%} exceeds the riskfree rate {rf:.4%} "
                      "(long-run nominal growth ceiling)", where("terminal_growth")))
    if wacc is not None and not WACC_BAND[0] <= wacc <= WACC_BAND[1]:
        out.append(_f("warning", "wacc_band", f"WACC {wacc:.4%} outside {WACC_BAND[0]:.0%}-{WACC_BAND[1]:.0%}; "
                      "re-check inputs and units", where("wacc")))
    pv_tv, ev = get("pv_terminal_value"), get("enterprise_value")
    if pv_tv is not None and ev:
        share = pv_tv / ev
        sev = "warning" if share > TV_SHARE_FLAG else "info"
        out.append(_f(sev, "tv_share_of_ev", f"PV of terminal value is {share:.1%} of EV"
                      + (" - value rests on the terminal assumptions; show the sensitivity" if sev == "warning" else ""),
                      where("pv_terminal_value")))
    periods, mid = get("discount_periods"), get("mid_year")
    observed, ambiguous = None, False
    if periods and len(periods) > 1:
        p1, p2 = periods[0], periods[1]
        # mid-year: p1 = s/2, p2 = s + 0.5 ; end-year: p1 = s, p2 = s + 1 (s = first-period fraction)
        is_mid = abs(p1 - (p2 - 0.5) / 2) < 1e-6
        is_end = abs(p1 - (p2 - 1)) < 1e-6
        if is_mid and is_end:  # [0.5, 1.5, ...]: full-year mid-year, or end-year after a half-year stub
            ambiguous = True
            observed = "end" if mid is False else "mid"
        elif is_mid:
            observed = "mid"
        elif is_end:
            observed = "end"
        steps = [b - a for a, b in zip(periods, periods[1:])]
        if any(abs(s - 1) > 1e-9 for s in steps):
            out.append(_f("warning", "discount_period_steps", f"discount periods do not step by 1: {periods}",
                          where("discount_periods")))
        out.append(_f("info", "discount_convention", f"discount periods {periods} read as "
                      f"{observed or 'unclear'}-year" + (" (ambiguous: also an end-year schedule after a "
                      "half-year stub; the mid-year switch decides)" if ambiguous else ""), where("discount_periods")))
    if mid is not None and observed is not None and not ambiguous and (observed == "mid") != mid:
        out.append(_f("critical", "mid_year_mismatch", f"mid-year switch is {'on' if mid else 'off'} but the discount "
                      f"periods are {observed}-year", where("mid_year")))
    tv = get("terminal_value")
    method = tv_method or ("exit" if get("exit_multiple") is not None and g is None else
                           "gordon" if g is not None else None)
    if tv and pv_tv and wacc and tv > 0 and pv_tv > 0 and wacc > 0:
        n = math.log(tv / pv_tv) / math.log(1 + wacc)
        last = periods[-1] if periods else None
        end_n = None if last is None or observed is None else (last + 0.5 if observed == "mid" else last)
        detail = f"TV discounted with exponent {n:.3f}"
        if end_n is not None:
            matches_end, matches_mid = abs(n - end_n) < 0.02, abs(n - (end_n - 0.5)) < 0.02
            if method == "exit" and matches_mid:
                out.append(_f("warning", "tv_timing", detail + f" (= N - 0.5 with N = {end_n:g}); an exit-multiple TV is "
                              "a value at the end of year N, so discount it at N", where("pv_terminal_value")))
            elif not (matches_end or matches_mid):
                out.append(_f("critical", "tv_timing", detail + f", which matches neither N = {end_n:g} nor N - 0.5: "
                              "TV is discounted with the wrong period (or not discounted)", where("pv_terminal_value")))
            else:
                out.append(_f("info", "tv_timing", detail + f" = {'N' if matches_end else 'N - 0.5'}; state this "
                              "convention next to the output"))
        else:
            out.append(_f("info", "tv_timing", detail + "; no discount-period row found to compare with"))
    wanted = ("wacc", "terminal_growth", "enterprise_value", "pv_terminal_value", "discount_periods")
    missing = [k for k in wanted if k not in found and not (k == "terminal_growth" and method == "exit")]
    if missing:
        out.append(_f("info", "inputs_not_found", "not located by label: " + ", ".join(missing)))
    return out


def summarise(findings):
    c = {s: sum(1 for f in findings if f["severity"] == s) for s in ("critical", "warning", "info")}
    c["overall"] = "Major issues" if c["critical"] else ("Minor issues" if c["warning"] else "Clean")
    return c


def main(argv=None) -> int:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    src = p.add_mutually_exclusive_group(required=True)
    src.add_argument("--xlsx")
    src.add_argument("--json")
    p.add_argument("--sheet", action="append")
    p.add_argument("--riskfree", type=float, help="override the riskfree rate used in the growth ceiling check")
    p.add_argument("--tv-method", choices=("gordon", "exit"))
    p.add_argument("--format", choices=("json", "table"), default="json")
    a = p.parse_args(argv)
    if a.xlsx:
        found, findings = read_xlsx(a.xlsx, a.sheet)
    else:
        with open(a.json, encoding="utf-8") as fh:
            data = json.load(fh)
        found = {k: (v, "json") for k, v in data.items() if k != "tv_method" and v is not None}
        if "mid_year" in found:
            found["mid_year"] = (to_flag(found["mid_year"][0]), "json")
        a.tv_method = a.tv_method or data.get("tv_method")
        findings = []
    findings += validate(found, a.riskfree, a.tv_method)
    s = summarise(findings)
    if a.format == "table":
        print(f"Overall: {s['overall']} - {s['critical']} critical, {s['warning']} warnings, {s['info']} info\n")
        print("| # | Severity | Check | Where | Detail |\n|---|---|---|---|---|")
        order = {"critical": 0, "warning": 1, "info": 2}
        for i, f in enumerate(sorted(findings, key=lambda f: order[f["severity"]]), 1):
            print(f"| {i} | {f['severity']} | {f['check']} | {f['where']} | {f['detail']} |")
    else:
        print(json.dumps({"summary": s, "findings": findings}, indent=1))
    return 1 if s["critical"] else 0


if __name__ == "__main__":
    sys.exit(main())
